"""
admin_app/views.py

Panel admin TravelHub. Semua view dilindungi @admin_required
(user harus login dan berstatus admin / superuser).
"""
import io
import re
from datetime import timedelta
from functools import wraps

from django.contrib import messages
from django.contrib.auth.views import redirect_to_login
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, ExpressionWrapper, F, IntegerField, Q, Sum
from django.db.models.functions import TruncDate
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST

from customer_app.models import Booking, Jadwal, Paket, Peserta, Tiket

from .forms import ItineraryFormSet, JadwalForm, PaketForm


# =========================================================
# BANTUAN
# =========================================================

def rp(angka):
    """12500000 -> 'Rp 12.500.000'"""
    return f"Rp {int(angka or 0):,}".replace(',', '.')


def admin_required(view):
    @wraps(view)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect_to_login(request.get_full_path())
        if not (request.user.is_superuser or request.user.role == 'admin'):
            messages.error(request, "Halaman ini khusus admin.")
            return redirect('home')
        return view(request, *args, **kwargs)
    return wrapper


def _paginate(request, queryset, per_page=10):
    page_obj = Paginator(queryset, per_page).get_page(request.GET.get('page'))
    params = request.GET.copy()
    params.pop('page', None)
    return page_obj, params.urlencode()


# total harga = harga paket x jumlah peserta (dihitung di database)
TOTAL_EXPR = ExpressionWrapper(
    F('jadwal__paket__harga') * F('jumlah_peserta'), output_field=IntegerField()
)

STATUS_BAYAR = ('paid', 'completed')   # booking yang dihitung sebagai pendapatan


# =========================================================
# 1. DASHBOARD
# =========================================================

@admin_required
def dashboard(request):
    hari_ini = timezone.localdate()

    pendapatan = (Booking.objects.filter(status__in=STATUS_BAYAR)
                  .aggregate(t=Sum(TOTAL_EXPR))['t'] or 0)
    paket_aktif = (Paket.objects
                   .filter(jadwal__tanggal_berangkat__gte=hari_ini).distinct().count())

    # ---- data grafik tren pemesanan ----
    rentang = 30 if request.GET.get('rentang') == '30' else 7
    mulai = hari_ini - timedelta(days=rentang - 1)
    per_hari = dict(
        Booking.objects.filter(created_at__date__gte=mulai)
        .exclude(status='cancelled')
        .annotate(tgl=TruncDate('created_at'))
        .values_list('tgl').annotate(n=Count('id'))
    )
    hari = [mulai + timedelta(days=i) for i in range(rentang)]
    nama_hari = ['Sen', 'Sel', 'Rab', 'Kam', 'Jum', 'Sab', 'Min']
    if rentang == 7:
        labels = [nama_hari[d.weekday()] for d in hari]
    else:
        labels = [f"{d.day}/{d.month}" for d in hari]
    data = [per_hari.get(d, 0) for d in hari]

    # jadwal yang kuotanya hampir habis
    hampir_penuh = (Jadwal.objects
                    .filter(tanggal_berangkat__gte=hari_ini, kuota_tersisa__lte=5)
                    .select_related('paket').order_by('kuota_tersisa')[:5])
    terbaru = (Booking.objects.select_related('user', 'jadwal__paket')
               .order_by('-created_at')[:5])

    return render(request, 'panel/dashboard.html', {
        'pendapatan': rp(pendapatan),
        'total_pesanan': Booking.objects.count(),
        'paket_aktif': paket_aktif,
        'belum_bayar': Booking.objects.filter(status='unpaid').count(),
        'rentang': rentang,
        'chart_labels': labels,
        'chart_data': data,
        'hampir_penuh': hampir_penuh,
        'terbaru': terbaru,
    })


# =========================================================
# 2. KELOLA PAKET
# =========================================================

@admin_required
def paket_list(request):
    q = request.GET.get('q', '').strip()
    paket = (Paket.objects.annotate(jml_jadwal=Count('jadwal', distinct=True))
             .order_by('-created_at'))
    if q:
        paket = paket.filter(Q(nama__icontains=q) | Q(destinasi__icontains=q))
    page_obj, qs = _paginate(request, paket)
    return render(request, 'panel/paket_list.html', {
        'page_obj': page_obj, 'query_string': qs, 'q': q,
    })


@admin_required
def paket_form(request, paket_id=None):
    paket = get_object_or_404(Paket, id=paket_id) if paket_id else None
    if request.method == 'POST':
        form = PaketForm(request.POST, request.FILES, instance=paket)
        formset = ItineraryFormSet(request.POST, instance=paket or Paket())
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                paket = form.save()
                formset.instance = paket
                formset.save()
            messages.success(request, f"Paket \"{paket.nama}\" berhasil disimpan.")
            return redirect('panel_paket')
    else:
        form = PaketForm(instance=paket)
        formset = ItineraryFormSet(instance=paket or Paket())
    return render(request, 'panel/paket_form.html', {
        'form': form, 'formset': formset, 'paket': paket,
    })


@admin_required
@require_POST
def paket_hapus(request, paket_id):
    paket = get_object_or_404(Paket, id=paket_id)
    # hapus paket = hapus jadwal & booking (CASCADE), jadi diblokir kalau sudah ada pesanan
    if Booking.objects.filter(jadwal__paket=paket).exists():
        messages.error(request, f"Paket \"{paket.nama}\" tidak bisa dihapus karena sudah punya pesanan.")
    else:
        paket.delete()
        messages.success(request, f"Paket \"{paket.nama}\" berhasil dihapus.")
    return redirect('panel_paket')


# =========================================================
# 3. JADWAL & KUOTA
# =========================================================

@admin_required
def jadwal_list(request):
    paket_id = request.GET.get('paket', '')
    jadwal = (Jadwal.objects.select_related('paket')
              .annotate(terisi=F('kuota_total') - F('kuota_tersisa'))
              .order_by('tanggal_berangkat'))
    if paket_id.isdigit():
        jadwal = jadwal.filter(paket_id=paket_id)
    page_obj, qs = _paginate(request, jadwal)
    return render(request, 'panel/jadwal_list.html', {
        'page_obj': page_obj, 'query_string': qs,
        'paket_list': Paket.objects.order_by('nama'),
        'paket_dipilih': paket_id,
        'hari_ini': timezone.localdate(),
    })


@admin_required
def jadwal_form(request, jadwal_id=None):
    jadwal = get_object_or_404(Jadwal, id=jadwal_id) if jadwal_id else None
    if request.method == 'POST':
        form = JadwalForm(request.POST, instance=jadwal)
        if form.is_valid():
            form.save()
            messages.success(request, "Jadwal berhasil disimpan.")
            return redirect('panel_jadwal')
    else:
        form = JadwalForm(instance=jadwal)
    return render(request, 'panel/jadwal_form.html', {'form': form, 'jadwal': jadwal})


@admin_required
@require_POST
def jadwal_hapus(request, jadwal_id):
    jadwal = get_object_or_404(Jadwal, id=jadwal_id)
    if jadwal.bookings.exists():
        messages.error(request, "Jadwal tidak bisa dihapus karena sudah ada pesanan.")
    else:
        jadwal.delete()
        messages.success(request, "Jadwal berhasil dihapus.")
    return redirect('panel_jadwal')


# =========================================================
# 4. MONITORING PESANAN
# =========================================================

@admin_required
def pesanan_list(request):
    status = request.GET.get('status', '')
    q = request.GET.get('q', '').strip()

    pesanan = (Booking.objects.select_related('user', 'jadwal__paket')
               .annotate(total=TOTAL_EXPR).order_by('-created_at'))
    if status in dict(Booking.STATUS_CHOICES):
        pesanan = pesanan.filter(status=status)
    if q:
        pesanan = pesanan.filter(
            Q(kode_booking__icontains=q) | Q(user__username__icontains=q) |
            Q(user__first_name__icontains=q) | Q(user__email__icontains=q) |
            Q(jadwal__paket__nama__icontains=q)
        )
    page_obj, qs = _paginate(request, pesanan, per_page=15)
    return render(request, 'panel/pesanan_list.html', {
        'page_obj': page_obj, 'query_string': qs, 'q': q, 'status': status,
        'status_choices': Booking.STATUS_CHOICES,
    })


@admin_required
def pesanan_detail(request, booking_id):
    booking = get_object_or_404(
        Booking.objects.select_related('user', 'jadwal__paket'), id=booking_id)
    return render(request, 'panel/pesanan_detail.html', {
        'b': booking,
        'peserta_list': booking.peserta.all(),
        'tiket': Tiket.objects.filter(booking=booking).first(),
    })


@admin_required
@require_POST
def pesanan_aksi(request, booking_id):
    aksi = request.POST.get('aksi')
    with transaction.atomic():
        booking = get_object_or_404(
            Booking.objects.select_for_update().select_related('jadwal'), id=booking_id)

        if aksi == 'bayar' and booking.status == 'unpaid':
            booking.konfirmasi_pembayaran()          # sekaligus menerbitkan tiket
            messages.success(request, f"{booking.kode_booking} dikonfirmasi lunas.")
        elif aksi == 'selesai' and booking.status == 'paid':
            booking.status = 'completed'
            booking.save()
            messages.success(request, f"{booking.kode_booking} ditandai selesai.")
        elif aksi == 'batal' and booking.status in ('unpaid', 'paid'):
            booking.status = 'cancelled'
            booking.save()
            booking.jadwal.tambah_kuota(booking.jumlah_peserta)   # kursi dikembalikan
            messages.success(request, f"{booking.kode_booking} dibatalkan, kuota dikembalikan.")
        else:
            messages.error(request, "Aksi tidak bisa dilakukan pada status pesanan saat ini.")
    return redirect('panel_pesanan_detail', booking_id=booking_id)


# =========================================================
# 5. VERIFIKASI TIKET (QR / KODE BOOKING)
# =========================================================

KODE_POLA = re.compile(r'TH-\d{4}-\d+', re.I)


@admin_required
def verifikasi(request):
    # POST = tombol "Check-in", GET = pencarian kode (dari input manual / hasil scan QR)
    kode = (request.POST.get('kode') or request.GET.get('kode') or '').strip()
    cocok = KODE_POLA.search(kode)          # isi QR bisa berupa teks/URL, ambil kodenya saja
    if cocok:
        kode = cocok.group(0)
    kode = kode.upper()

    booking = tiket = None
    if kode:
        booking = (Booking.objects.select_related('user', 'jadwal__paket')
                   .filter(kode_booking__iexact=kode).first())
        if not booking:
            messages.error(request, f"Kode \"{kode}\" tidak ditemukan.")
        elif booking.status == 'paid':
            tiket, _ = Tiket.objects.get_or_create(booking=booking)

    if request.method == 'POST' and tiket:
        if tiket.sudah_diverifikasi:
            messages.info(request, "Tiket ini sudah pernah check-in.")
        else:
            tiket.verifikasi()
            messages.success(request, f"Check-in berhasil untuk {booking.kode_booking}.")
        return redirect(f"{request.path}?kode={booking.kode_booking}")

    return render(request, 'panel/verifikasi.html', {
        'kode': kode, 'b': booking, 'tiket': tiket,
        'peserta_list': booking.peserta.all() if booking else [],
        'hari_ini': timezone.localdate(),
    })


# =========================================================
# 6. EKSPOR LAPORAN
# =========================================================

JENIS_LAPORAN = {
    'penjualan': 'Laporan Penjualan Tiket',
    'peserta': 'Laporan Data Peserta',
    'pembayaran': 'Laporan Status Pembayaran',
}


def _nama(user):
    return user.get_full_name() or user.username


def _tgl(d):
    return d.strftime('%d-%m-%Y') if d else ''


def _susun_laporan(jenis, dari, sampai):
    """Return (header, baris, indeks_kolom_uang)."""
    booking = Booking.objects.select_related('user', 'jadwal__paket').order_by('created_at')
    if dari:
        booking = booking.filter(created_at__date__gte=dari)
    if sampai:
        booking = booking.filter(created_at__date__lte=sampai)

    if jenis == 'penjualan':
        header = ['Kode Booking', 'Tgl Pesan', 'Pelanggan', 'Paket', 'Tgl Berangkat',
                  'Jumlah', 'Harga/Orang', 'Total']
        baris = [[b.kode_booking, _tgl(b.created_at), _nama(b.user), b.jadwal.paket.nama,
                  _tgl(b.jadwal.tanggal_berangkat), b.jumlah_peserta,
                  b.jadwal.paket.harga, b.total_harga()]
                 for b in booking.filter(status__in=STATUS_BAYAR)]
        if baris:
            baris.append(['TOTAL', '', '', '', '', sum(r[5] for r in baris), '',
                          sum(r[7] for r in baris)])
        return header, baris, {6, 7}

    if jenis == 'peserta':
        peserta = (Peserta.objects.filter(booking__in=booking.exclude(status='cancelled'))
                   .select_related('booking__jadwal__paket').order_by('booking__created_at', 'id'))
        header = ['Kode Booking', 'Nama Peserta', 'No Identitas', 'Tgl Lahir', 'JK',
                  'Kontak Darurat', 'Paket', 'Tgl Berangkat', 'Status']
        baris = [[p.booking.kode_booking, p.nama_lengkap, p.no_identitas,
                  _tgl(p.tanggal_lahir), p.jenis_kelamin,
                  f"{p.kontak_darurat_nama} {p.kontak_darurat_hp}".strip(),
                  p.booking.jadwal.paket.nama,
                  _tgl(p.booking.jadwal.tanggal_berangkat), p.booking.get_status_display()]
                 for p in peserta]
        return header, baris, set()

    header = ['Kode Booking', 'Tgl Pesan', 'Pelanggan', 'Email', 'Paket', 'Total',
              'Metode', 'Status']
    baris = [[b.kode_booking, _tgl(b.created_at), _nama(b.user), b.user.email,
              b.jadwal.paket.nama, b.total_bayar(), b.metode_pembayaran.upper() or '-',
              b.get_status_display()]
             for b in booking]
    return header, baris, {5}


def _xlsx(judul, periode, header, baris, uang):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Laporan'
    ws['A1'] = judul
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = periode
    ws.append([])
    ws.append(header)
    for c in ws[4]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='4F46E5')
        c.alignment = Alignment(horizontal='center')
    for r in baris:
        ws.append(r)
    for row in ws.iter_rows(min_row=5):
        for i, c in enumerate(row):
            if i in uang and isinstance(c.value, int):
                c.number_format = '#,##0'
    if baris and baris[-1][0] == 'TOTAL':
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for i, h in enumerate(header, 1):
        panjang = max([len(str(h))] + [len(str(r[i - 1])) for r in baris])
        ws.column_dimensions[get_column_letter(i)].width = min(panjang + 3, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(judul, periode, header, baris, uang):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24,
                            topMargin=28, bottomMargin=28)
    gaya = getSampleStyleSheet()
    isi = [[(rp(v) if (i in uang and isinstance(v, int)) else v) for i, v in enumerate(r)]
           for r in baris]
    tabel = Table([header] + isi, repeatRows=1)
    style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]
    if baris and baris[-1][0] == 'TOTAL':
        style.append(('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'))
    tabel.setStyle(TableStyle(style))
    doc.build([Paragraph(judul, gaya['Title']), Paragraph(periode, gaya['Normal']),
               Spacer(1, 12), tabel])
    return buf.getvalue()


@admin_required
def laporan(request):
    return render(request, 'panel/laporan.html', {
        'jenis_laporan': JENIS_LAPORAN,
        'hari_ini': timezone.localdate(),
    })


@admin_required
def laporan_ekspor(request):
    jenis = request.GET.get('jenis')
    fmt = request.GET.get('format')
    dari = parse_date(request.GET.get('dari') or '')
    sampai = parse_date(request.GET.get('sampai') or '')

    if jenis not in JENIS_LAPORAN or fmt not in ('xlsx', 'pdf'):
        messages.error(request, "Pilihan laporan tidak valid.")
        return redirect('panel_laporan')
    if dari and sampai and dari > sampai:
        messages.error(request, "Tanggal mulai tidak boleh lebih besar dari tanggal akhir.")
        return redirect('panel_laporan')

    header, baris, uang = _susun_laporan(jenis, dari, sampai)
    if not baris:
        messages.error(request, "Tidak ada data pada periode tersebut.")
        return redirect('panel_laporan')

    judul = JENIS_LAPORAN[jenis]
    periode = f"Periode: {_tgl(dari) or 'awal'} s/d {_tgl(sampai) or 'sekarang'}"
    try:
        if fmt == 'xlsx':
            isi = _xlsx(judul, periode, header, baris, uang)
            tipe = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        else:
            isi = _pdf(judul, periode, header, baris, uang)
            tipe = 'application/pdf'
    except ImportError:
        lib = 'openpyxl' if fmt == 'xlsx' else 'reportlab'
        messages.error(request, f"Library {lib} belum terpasang. Jalankan: pip install {lib}")
        return redirect('panel_laporan')

    nama_file = f"{jenis}_{timezone.localdate():%Y%m%d}.{fmt}"
    resp = HttpResponse(isi, content_type=tipe)
    resp['Content-Disposition'] = f'attachment; filename="{nama_file}"'
    return resp
