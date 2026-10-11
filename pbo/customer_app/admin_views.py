import io
from functools import wraps

from django.contrib import messages
from django.contrib.auth import get_user_model, logout as auth_logout
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db.models import Count, ExpressionWrapper, F, IntegerField, Sum
from django.db.models.functions import Coalesce, TruncMonth
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from .admin_forms import JadwalAdminForm, LaporanForm, PaketAdminForm, VerifikasiForm
from .models import Booking, Jadwal, Paket, Tiket

User = get_user_model()

# Booking berstatus 'paid' dan 'completed' (perjalanan selesai) sama-sama dihitung sebagai pendapatan.
# Kalau hanya ingin 'paid' saja, ubah menjadi ('paid',).
STATUS_LUNAS = ("paid", "completed")

BULAN_ID = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]

# Booking tidak punya field total_harga (hanya method), jadi dihitung di database:
# jumlah_peserta x harga paket.
TOTAL_HARGA = ExpressionWrapper(
    F("jumlah_peserta") * F("jadwal__paket__harga"), output_field=IntegerField()
)

MENU = [
    ("dashboard", "Dashboard", "layout-dashboard", "admin_dashboard"),
    ("paket", "Paket Wisata", "package", "admin_paket"),
    ("jadwal", "Jadwal & Kuota", "calendar", "admin_jadwal"),
    ("pesanan", "Pesanan", "clipboard-list", "admin_pesanan"),
    ("verifikasi", "Verifikasi QR", "qr-code", "admin_verifikasi"),
    ("laporan", "Ekspor Laporan", "file-down", "admin_laporan"),
]


# ---------------------------------------------------------------------------
# Helper
# ---------------------------------------------------------------------------
def admin_required(view_func):
    """Harus login dan role 'admin' (atau superuser)."""

    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        user = request.user
        if not (user.is_superuser or user.role == "admin"):
            raise PermissionDenied
        return view_func(request, *args, **kwargs)

    return login_required(_wrapped)


def _render(request, template, active, **context):
    context.update(menu=MENU, active_menu=active)
    return render(request, f"admin_panel/{template}", context)


def _rupiah(value):
    return "Rp " + f"{int(value or 0):,}".replace(",", ".")


def _nama(user):
    # RegisterForm menyimpan nama lengkap di first_name
    return user.get_full_name() or user.username


@require_POST
def admin_keluar(request):
    auth_logout(request)
    return redirect("home")


# ---------------------------------------------------------------------------
# 1. DASHBOARD
# ---------------------------------------------------------------------------
@admin_required
def admin_dashboard(request):
    now = timezone.localtime()

    # --- kartu statistik ---
    pendapatan_bulan_ini = Booking.objects.filter(
        status__in=STATUS_LUNAS, created_at__year=now.year, created_at__month=now.month
    ).aggregate(t=Coalesce(Sum(TOTAL_HARGA), 0))["t"]

    kuota = Jadwal.objects.aggregate(
        total=Coalesce(Sum("kuota_total"), 0), sisa=Coalesce(Sum("kuota_tersisa"), 0)
    )
    terisi = kuota["total"] - kuota["sisa"]
    okupansi_rata = round(terisi / kuota["total"] * 100, 1) if kuota["total"] else 0

    total_pesanan = Booking.objects.count()
    pelanggan_aktif = User.objects.filter(role="customer", is_active=True).count()

    # --- grafik garis: pendapatan 12 bulan terakhir ---
    bulan_list = []
    y, m = now.year, now.month
    for _ in range(12):
        bulan_list.append((y, m))
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    bulan_list.reverse()

    awal = now.replace(year=bulan_list[0][0], month=bulan_list[0][1], day=1,
                       hour=0, minute=0, second=0, microsecond=0)
    per_bulan = {
        (r["bulan"].year, r["bulan"].month): r["total"]
        for r in Booking.objects.filter(status__in=STATUS_LUNAS, created_at__gte=awal)
        .annotate(bulan=TruncMonth("created_at"))
        .values("bulan")
        .annotate(total=Sum(TOTAL_HARGA))
    }
    chart_pendapatan = {
        "labels": [f"{BULAN_ID[b - 1]} {t}" for t, b in bulan_list],
        "values": [per_bulan.get(k, 0) for k in bulan_list],
    }

    # --- grafik batang: okupansi per paket ---
    agregat = {
        r["paket_id"]: r
        for r in Jadwal.objects.values("paket_id")
        .annotate(total=Sum("kuota_total"), sisa=Sum("kuota_tersisa"))
        .order_by()
    }
    labels, values = [], []
    for p in Paket.objects.order_by("nama"):
        r = agregat.get(p.id)
        persen = round((r["total"] - r["sisa"]) / r["total"] * 100, 1) if r and r["total"] else 0
        labels.append(p.nama)
        values.append(persen)
    chart_okupansi = {"labels": labels, "values": values}

    # --- diagram donat: pesanan per status ---
    jumlah = {r["status"]: r["n"] for r in Booking.objects.values("status").annotate(n=Count("id")).order_by()}
    chart_status = {
        "labels": ["Paid", "Unpaid", "Cancelled", "Completed"],
        "values": [jumlah.get("paid", 0), jumlah.get("unpaid", 0),
                   jumlah.get("cancelled", 0), jumlah.get("completed", 0)],
    }

    return _render(
        request, "dashboard.html", "dashboard",
        pendapatan_bulan_ini=pendapatan_bulan_ini,
        okupansi_rata=okupansi_rata,
        total_pesanan=total_pesanan,
        pelanggan_aktif=pelanggan_aktif,
        chart_data={"pendapatan": chart_pendapatan, "okupansi": chart_okupansi, "status": chart_status},
    )


# ---------------------------------------------------------------------------
# 2. PAKET WISATA (CRUD)
# ---------------------------------------------------------------------------
@admin_required
def paket_list(request):
    return _render(request, "paket_list.html", "paket", pakets=Paket.objects.all())


@admin_required
def paket_tambah(request):
    form = PaketAdminForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Paket wisata berhasil ditambahkan.")
        return redirect("admin_paket")
    return _render(request, "paket_form.html", "paket", form=form, judul="Tambah Paket Wisata")


@admin_required
def paket_edit(request, pk):
    paket = get_object_or_404(Paket, pk=pk)
    form = PaketAdminForm(request.POST or None, request.FILES or None, instance=paket)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Paket wisata berhasil diperbarui.")
        return redirect("admin_paket")
    return _render(request, "paket_form.html", "paket", form=form, paket=paket, judul=f"Edit Paket: {paket.nama}")


@admin_required
@require_POST
def paket_hapus(request, pk):
    paket = get_object_or_404(Paket, pk=pk)
    # Booking -> Jadwal -> Paket semuanya CASCADE: menghapus paket yang sudah punya pesanan
    # akan ikut menghapus seluruh riwayat pesanan pelanggan. Karena itu diblok di sini.
    if Booking.objects.filter(jadwal__paket=paket).exists():
        messages.error(request, f"Paket '{paket.nama}' tidak bisa dihapus karena sudah memiliki pesanan.")
    else:
        nama = paket.nama
        paket.delete()
        messages.success(request, f"Paket '{nama}' berhasil dihapus.")
    return redirect("admin_paket")


# ---------------------------------------------------------------------------
# 3. JADWAL & KUOTA
# ---------------------------------------------------------------------------
@admin_required
def jadwal_list(request):
    if request.method == "POST":
        form = JadwalAdminForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Jadwal keberangkatan berhasil ditambahkan.")
            return redirect("admin_jadwal")
    else:
        form = JadwalAdminForm()

    jadwals = list(Jadwal.objects.select_related("paket").order_by("tanggal_berangkat", "paket__nama"))
    for j in jadwals:
        j.terisi = j.kuota_total - j.kuota_tersisa
        j.persen = round(j.terisi / j.kuota_total * 100) if j.kuota_total else 0

    return _render(request, "jadwal.html", "jadwal", jadwals=jadwals, form=form)


# ---------------------------------------------------------------------------
# 4. PESANAN
# ---------------------------------------------------------------------------
@admin_required
def pesanan_list(request):
    status = request.GET.get("status", "").lower()
    if status not in ("unpaid", "paid", "cancelled"):
        status = ""

    qs = (Booking.objects.select_related("user", "jadwal__paket")
          .annotate(total=TOTAL_HARGA).order_by("-created_at"))
    if status:
        qs = qs.filter(status=status)

    jumlah = {r["status"]: r["n"] for r in Booking.objects.values("status").annotate(n=Count("id")).order_by()}
    tabs = [
        ("", "Semua", sum(jumlah.values())),
        ("unpaid", "Unpaid", jumlah.get("unpaid", 0)),
        ("paid", "Paid", jumlah.get("paid", 0)),
        ("cancelled", "Cancelled", jumlah.get("cancelled", 0)),
    ]

    page_obj = Paginator(qs, 20).get_page(request.GET.get("page"))
    return _render(request, "pesanan.html", "pesanan", page_obj=page_obj, tabs=tabs, status_aktif=status)


# ---------------------------------------------------------------------------
# 5. VERIFIKASI TIKET (input manual kode booking)
# ---------------------------------------------------------------------------
BADGE = {
    "valid": "bg-emerald-100 text-emerald-700",
    "sudah_checkin": "bg-sky-100 text-sky-700",
    "belum_bayar": "bg-amber-100 text-amber-700",
    "dibatalkan": "bg-slate-200 text-slate-700",
    "tidak_ditemukan": "bg-red-100 text-red-700",
}


def _hasil(status, label, kode, booking=None, tiket=None):
    return {"status": status, "label": label, "kode": kode, "booking": booking,
            "tiket": tiket, "badge": BADGE[status]}


@admin_required
def verifikasi(request):
    hasil = None
    if request.method == "POST":
        form = VerifikasiForm(request.POST)
        if form.is_valid():
            kode = form.cleaned_data["kode"]
            booking = (Booking.objects.select_related("user", "jadwal__paket")
                       .annotate(total=TOTAL_HARGA).filter(kode_booking__iexact=kode).first())

            if booking is None:
                hasil = _hasil("tidak_ditemukan", "Tidak Ditemukan", kode)
            elif booking.status == "unpaid":
                hasil = _hasil("belum_bayar", "Belum Bayar", kode, booking)
            elif booking.status == "cancelled":
                hasil = _hasil("dibatalkan", "Dibatalkan", kode, booking)
            else:  # paid / completed -> punya tiket
                tiket, _ = Tiket.objects.get_or_create(booking=booking)
                if request.POST.get("aksi") == "checkin" and not tiket.sudah_diverifikasi:
                    tiket.verifikasi()
                    messages.success(request, f"Check-in {booking.kode_booking} berhasil dicatat.")
                if tiket.sudah_diverifikasi:
                    hasil = _hasil("sudah_checkin", "Sudah Check-in", kode, booking, tiket)
                else:
                    hasil = _hasil("valid", "Valid", kode, booking, tiket)
    else:
        form = VerifikasiForm()

    return _render(request, "verifikasi.html", "verifikasi", form=form, hasil=hasil)


# ---------------------------------------------------------------------------
# 6. EKSPOR LAPORAN
# ---------------------------------------------------------------------------
def _build_report(jenis, d1, d2):
    periode = f"Periode: {d1:%d-%m-%Y} s/d {d2:%d-%m-%Y}"

    if jenis == "pendapatan":
        qs = (Booking.objects.select_related("user", "jadwal__paket")
              .filter(status__in=STATUS_LUNAS, created_at__date__range=(d1, d2))
              .annotate(total=TOTAL_HARGA).order_by("created_at"))
        rows = [
            [b.kode_booking, timezone.localtime(b.created_at).strftime("%d-%m-%Y"), _nama(b.user),
             b.jadwal.paket.nama, b.jumlah_peserta, b.total]
            for b in qs
        ]
        return {
            "title": "Laporan Pendapatan", "sheet": "Pendapatan", "periode": periode,
            "headers": ["Kode Booking", "Tanggal", "Pelanggan", "Paket", "Peserta", "Total (Rp)"],
            "rows": rows, "footer": ["", "", "", "", "TOTAL", sum(r[5] for r in rows)],
            "currency_cols": [5],
        }

    if jenis == "pesanan":
        qs = (Booking.objects.select_related("user", "jadwal__paket")
              .filter(created_at__date__range=(d1, d2))
              .annotate(total=TOTAL_HARGA).order_by("created_at"))
        rows = [
            [b.kode_booking, timezone.localtime(b.created_at).strftime("%d-%m-%Y"), _nama(b.user),
             b.jadwal.paket.nama, b.jadwal.tanggal_berangkat.strftime("%d-%m-%Y"), b.total,
             b.get_status_display()]
            for b in qs
        ]
        return {
            "title": "Laporan Pesanan", "sheet": "Pesanan", "periode": periode,
            "headers": ["Kode Booking", "Tgl Pesan", "Pelanggan", "Paket", "Tgl Berangkat", "Total (Rp)", "Status"],
            "rows": rows, "footer": None, "currency_cols": [5],
        }

    # okupansi: jadwal yang berangkat dalam rentang tanggal
    qs = (Jadwal.objects.select_related("paket")
          .filter(tanggal_berangkat__range=(d1, d2)).order_by("tanggal_berangkat", "paket__nama"))
    rows = []
    for j in qs:
        terisi = j.kuota_total - j.kuota_tersisa
        persen = terisi / j.kuota_total * 100 if j.kuota_total else 0
        rows.append([j.paket.nama, j.tanggal_berangkat.strftime("%d-%m-%Y"), j.kuota_total, terisi,
                     f"{persen:.1f}%"])
    return {
        "title": "Laporan Okupansi Jadwal", "sheet": "Okupansi", "periode": periode,
        "headers": ["Paket", "Tgl Berangkat", "Kuota", "Terisi", "Okupansi"],
        "rows": rows, "footer": None, "currency_cols": [],
    }


def _export_excel(report, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = report["sheet"]
    ncol = len(report["headers"])

    ws.append([report["title"]])
    ws.append([report["periode"]])
    ws.append([])
    ws.append(report["headers"])
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(italic=True)

    header_row = 4
    for c in range(1, ncol + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A2C40")
        cell.alignment = Alignment(horizontal="center")

    for row in report["rows"]:
        ws.append(row)
    if report["footer"]:
        ws.append(report["footer"])
        for c in range(1, ncol + 1):
            ws.cell(row=ws.max_row, column=c).font = Font(bold=True)

    for col in report["currency_cols"]:
        for r in range(header_row + 1, ws.max_row + 1):
            ws.cell(row=r, column=col + 1).number_format = "#,##0"

    for c in range(1, ncol + 1):
        lebar = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(header_row, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(c)].width = min(max(lebar + 3, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def _export_pdf(report, filename):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def fmt(row):
        return [_rupiah(v) if i in report["currency_cols"] and isinstance(v, (int, float)) else str(v)
                for i, v in enumerate(row)]

    data = [report["headers"]] + [fmt(r) for r in report["rows"]]
    if report["footer"]:
        data.append(fmt(report["footer"]))
    if len(data) == 1:
        data.append(["Tidak ada data pada periode ini"] + [""] * (len(report["headers"]) - 1))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24,
                            topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    table = Table(data, repeatRows=1)
    gaya = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A2C40")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F7F9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if report["footer"]:
        gaya += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                 ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#DCE9EF"))]
    table.setStyle(TableStyle(gaya))

    doc.build([
        Paragraph(report["title"], styles["Title"]),
        Paragraph(report["periode"], styles["Normal"]),
        Spacer(1, 12),
        table,
    ])
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return response


@admin_required
def laporan(request):
    if request.method == "POST":
        form = LaporanForm(request.POST)
        if form.is_valid():
            d = form.cleaned_data
            report = _build_report(d["jenis"], d["tanggal_dari"], d["tanggal_sampai"])
            filename = f"laporan_{d['jenis']}_{d['tanggal_dari']:%Y%m%d}_{d['tanggal_sampai']:%Y%m%d}"
            if d["format"] == "excel":
                return _export_excel(report, filename)
            return _export_pdf(report, filename)
    else:
        hari_ini = timezone.localdate()
        form = LaporanForm(initial={"tanggal_dari": hari_ini.replace(day=1), "tanggal_sampai": hari_ini})

    return _render(request, "laporan.html", "laporan", form=form)